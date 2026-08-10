"""LibreOffice が無いので、pptx の座標を読んで matplotlib で見た目を再現する。

正確なレンダリングではないが、重なり・余白・行の詰まりを目で確かめられる。
"""
import os
import sys
import unicodedata
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.image as mpimg
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from mathtext import omath_of, linearize, eq_size

plt.rcParams["font.family"] = "sans-serif"
plt.rcParams["font.sans-serif"] = ["Hiragino Sans", "DejaVu Sans"]

EMU = 914400.0
W, H = 13.333, 7.5
HERE = os.path.dirname(os.path.abspath(__file__))
OUTD = os.path.join(HERE, "preview")
os.makedirs(OUTD, exist_ok=True)


def em(ch):
    return 1.0 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 0.52


def wrap(text, width_in, size_pt):
    cap = width_in * 72.0 / size_pt
    out, cur, acc = [], "", 0.0
    for ch in text:
        if ch == "\n":
            out.append(cur); cur, acc = "", 0.0; continue
        w = em(ch)
        if acc + w > cap:
            out.append(cur); cur, acc = ch, w
        else:
            cur += ch; acc += w
    out.append(cur)
    return out


def rgb(c):
    if c is None:
        return None
    return "#" + str(c)


def render(path):
    prs = Presentation(path)
    for idx, slide in enumerate(prs.slides, start=1):
        fig, ax = plt.subplots(figsize=(W, H))
        ax.set_position([0, 0, 1, 1])
        ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")
        bgc = "#FFFFFF"
        try:
            if slide.background.fill.type is not None and \
                    slide.background.fill.type == 1:
                bgc = rgb(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        ax.add_patch(Rectangle((0, 0), W, H, facecolor=bgc, edgecolor="#DDDDDD"))
        for sh in slide.shapes:
            x, y = sh.left / EMU, sh.top / EMU
            w, h = sh.width / EMU, sh.height / EMU
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = mpimg.imread(
                    os.path.join(HERE, "figs", "_tmp.png")
                    if False else sh.image.blob and _blob_path(sh, idx))
                ax.imshow(img, extent=(x, x + w, y + h, y), aspect="auto",
                          zorder=2)
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.TABLE:
                tbl = sh.table
                ys = y
                for r in tbl.rows:
                    rh = r.height / EMU
                    xs = x
                    for j, cell in enumerate(r.cells):
                        cw = tbl.columns[j].width / EMU
                        fc = "#FFFFFF"
                        try:
                            fc = rgb(cell.fill.fore_color.rgb)
                        except Exception:
                            pass
                        ax.add_patch(Rectangle((xs, ys), cw, rh, facecolor=fc,
                                               edgecolor="#CCCCCC", lw=0.6,
                                               zorder=2))
                        txt = cell.text
                        runs = [rr for p in cell.text_frame.paragraphs
                                for rr in p.runs]
                        sz = runs[0].font.size.pt if runs and runs[0].font.size else 12
                        col = "#222222"
                        if runs and runs[0].font.color and runs[0].font.color.type is not None:
                            try:
                                col = rgb(runs[0].font.color.rgb)
                            except Exception:
                                pass
                        lines = wrap(txt, cw - 0.18, sz)
                        for k, ln in enumerate(lines):
                            ax.text(xs + 0.10, ys + rh / 2
                                    - (len(lines) - 1) * sz * 1.3 / 144
                                    + k * sz * 1.3 / 72,
                                    ln, fontsize=sz, color=col, va="center",
                                    ha="left", zorder=3)
                        xs += cw
                    ys += rh
                continue
            # 図形（塗り）
            if sh.has_text_frame and sh.shape_type is not None and \
                    "TEXT_BOX" not in str(sh.shape_type):
                fc, ec = "#FFFFFF", None
                try:
                    fc = rgb(sh.fill.fore_color.rgb)
                except Exception:
                    fc = None
                try:
                    ec = rgb(sh.line.color.rgb)
                except Exception:
                    ec = None
                ax.add_patch(Rectangle((x, y), w, h,
                                       facecolor=fc or "none",
                                       edgecolor=ec or "none",
                                       lw=1.2, zorder=1))
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            cy = y
            for p in tf.paragraphs:
                runs = p.runs
                om = omath_of(p)
                if om is not None:
                    sz = eq_size(p)
                    col = "#1B3654"
                    bold = False
                    txt = linearize(om)
                elif not runs:
                    continue
                else:
                    sz = max((r.font.size.pt for r in runs if r.font.size),
                             default=18)
                    col = "#222222"
                    try:
                        col = rgb(runs[0].font.color.rgb)
                    except Exception:
                        pass
                    bold = bool(runs[0].font.bold)
                    txt = "".join(r.text for r in runs)
                ls = p.line_spacing or 1.32
                lines = wrap(txt, w, sz)
                align = str(p.alignment)
                for ln in lines:
                    if "CENTER" in align:
                        ax.text(x + w / 2, cy + sz * ls / 144, ln, fontsize=sz,
                                color=col, ha="center", va="center",
                                fontweight="bold" if bold else "normal", zorder=4)
                    elif "RIGHT" in align:
                        ax.text(x + w, cy + sz * ls / 144, ln, fontsize=sz,
                                color=col, ha="right", va="center",
                                fontweight="bold" if bold else "normal", zorder=4)
                    else:
                        ax.text(x, cy + sz * ls / 144, ln, fontsize=sz,
                                color=col, ha="left", va="center",
                                fontweight="bold" if bold else "normal", zorder=4)
                    cy += sz * ls / 72
                cy += (p.space_after.pt if p.space_after else 0) / 72
        out = os.path.join(OUTD, f"p{idx:02d}.png")
        fig.savefig(out, dpi=80,
                    facecolor=bgc)
        plt.close(fig)
    print("preview:", OUTD)


_cache = {}


def _blob_path(sh, idx):
    key = sh.image.sha1
    if key not in _cache:
        p = os.path.join(OUTD, f"_img_{key[:8]}.png")
        with open(p, "wb") as f:
            f.write(sh.image.blob)
        _cache[key] = p
    return _cache[key]


if __name__ == "__main__":
    render(sys.argv[1])
