"""PowerPoint の「挿入 → 数式」と同じ形式（OMML）を組み立てる。

python-pptx は数式を扱えないので、a:p の中へ
  <a14:m><m:oMathPara><m:oMath> ... </m:oMath></m:oMathPara></a14:m>
を直接差し込む。PowerPoint 側では通常の数式として編集できる。

使い方:
    from omml import Eq, add_equation
    add_equation(text_frame, Eq.seq(Eq.i("x"), Eq.o("="), Eq.n("2")), size=16)
"""
from pptx.oxml import parse_xml

M_URI = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A_URI = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14_URI = "http://schemas.microsoft.com/office/drawing/2010/main"
MC_URI = "http://schemas.openxmlformats.org/markup-compatibility/2006"

_MATH_FONT = ('<a:latin typeface="Cambria Math" panose="02040503050406030204" '
              'pitchFamily="18" charset="0"/>')


def _esc(t):
    return (t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def _rpr(sz, italic, color):
    return (f'<a:rPr lang="en-US" sz="{int(sz * 100)}" b="0" '
            f'i="{1 if italic else 0}" dirty="0">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'{_MATH_FONT}</a:rPr>')


def _ctrl(sz, color):
    return f'<m:ctrlPr>{_rpr(sz, True, color)}</m:ctrlPr>'


class _Node:
    """size / color は組み立ての最後にまとめて差し込む（テンプレート方式）。"""

    def __init__(self, tmpl):
        self.tmpl = tmpl

    def render(self, sz, color):
        return self.tmpl.format(RPR_I=_rpr(sz, True, color),
                                RPR_U=_rpr(sz, False, color),
                                CTRL=_ctrl(sz, color))


class Eq:
    """数式の部品。i=斜体の変数, u=立体の文字, o=演算子, n=数字。"""

    @staticmethod
    def _run(text, italic):
        key = "{RPR_I}" if italic else "{RPR_U}"
        return _Node("<m:r>" + key
                     + f'<m:t xml:space="preserve">{_esc(text)}</m:t></m:r>')

    @staticmethod
    def i(text):
        return Eq._run(text, True)

    @staticmethod
    def u(text):
        return Eq._run(text, False)

    @staticmethod
    def o(text):
        return Eq._run(f" {text} ", False)

    @staticmethod
    def n(text):
        return Eq._run(text, False)

    @staticmethod
    def sp():
        return Eq._run(" ", False)

    @staticmethod
    def seq(*parts):
        return _Node("".join(p.tmpl for p in parts))

    @staticmethod
    def sub(base, sub):
        return _Node("<m:sSub><m:sSubPr>{CTRL}</m:sSubPr>"
                     f"<m:e>{base.tmpl}</m:e><m:sub>{sub.tmpl}</m:sub>"
                     "</m:sSub>")

    @staticmethod
    def sup(base, sup):
        return _Node("<m:sSup><m:sSupPr>{CTRL}</m:sSupPr>"
                     f"<m:e>{base.tmpl}</m:e><m:sup>{sup.tmpl}</m:sup>"
                     "</m:sSup>")

    @staticmethod
    def subsup(base, sub, sup):
        return _Node("<m:sSubSup><m:sSubSupPr>{CTRL}</m:sSubSupPr>"
                     f"<m:e>{base.tmpl}</m:e><m:sub>{sub.tmpl}</m:sub>"
                     f"<m:sup>{sup.tmpl}</m:sup></m:sSubSup>")

    @staticmethod
    def frac(num, den):
        return _Node("<m:f><m:fPr>{CTRL}</m:fPr>"
                     f"<m:num>{num.tmpl}</m:num><m:den>{den.tmpl}</m:den>"
                     "</m:f>")

    @staticmethod
    def d(inner, beg="(", end=")"):
        return _Node(f'<m:d><m:dPr><m:begChr m:val="{beg}"/>'
                     f'<m:endChr m:val="{end}"/>' + "{CTRL}</m:dPr>"
                     f"<m:e>{inner.tmpl}</m:e></m:d>")

    @staticmethod
    def nary(chr_, sub, body, sup=None, limloc="subSup"):
        sup_xml = f"<m:sup>{sup.tmpl}</m:sup>" if sup is not None else "<m:sup/>"
        hide = "" if sup is not None else '<m:supHide m:val="1"/>'
        return _Node(f'<m:nary><m:naryPr><m:chr m:val="{chr_}"/>'
                     f'<m:limLoc m:val="{limloc}"/>{hide}' + "{CTRL}</m:naryPr>"
                     f"<m:sub>{sub.tmpl}</m:sub>{sup_xml}"
                     f"<m:e>{body.tmpl}</m:e></m:nary>")

    @staticmethod
    def bar(inner):
        return _Node('<m:bar><m:barPr><m:pos m:val="top"/>' + "{CTRL}</m:barPr>"
                     f"<m:e>{inner.tmpl}</m:e></m:bar>")

    @staticmethod
    def mat(rows):
        """rows: list[list[_Node]] — 行列。"""
        n_col = len(rows[0])
        cols = ('<m:mcs><m:mc><m:mcPr>'
                f'<m:count m:val="{n_col}"/>'
                '<m:mcJc m:val="center"/></m:mcPr></m:mc></m:mcs>')
        body = "".join(
            "<m:mr>" + "".join(f"<m:e>{c.tmpl}</m:e>" for c in row) + "</m:mr>"
            for row in rows)
        return _Node(f"<m:m><m:mPr>{cols}" + "{CTRL}</m:mPr>" + body + "</m:m>")


def equation_xml(node, size=16.0, color="222222", align="left"):
    inner = node.render(size, color)
    return (f'<a14:m xmlns:a14="{A14_URI}" xmlns:a="{A_URI}">'
            f'<m:oMathPara xmlns:m="{M_URI}">'
            f'<m:oMathParaPr><m:jc m:val="{align}"/></m:oMathParaPr>'
            f"<m:oMath>{inner}</m:oMath>"
            "</m:oMathPara></a14:m>")


def add_equation(text_frame, node, size=16.0, color="222222", align="left",
                 space_before=0, space_after=6, first=False):
    """text_frame に数式だけの段落を1つ足す。first=True なら先頭段落を使う。"""
    if first and len(text_frame.paragraphs) == 1 and not text_frame.paragraphs[0].runs:
        p = text_frame.paragraphs[0]
    else:
        p = text_frame.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    if space_before:
        pPr.append(parse_xml(
            f'<a:spcBef xmlns:a="{A_URI}"><a:spcPts val="{int(space_before*100)}"/>'
            "</a:spcBef>"))
    if space_after:
        pPr.append(parse_xml(
            f'<a:spcAft xmlns:a="{A_URI}"><a:spcPts val="{int(space_after*100)}"/>'
            "</a:spcAft>"))
    p._p.append(parse_xml(equation_xml(node, size, color, align)))
    p._p.append(parse_xml(
        f'<a:endParaRPr xmlns:a="{A_URI}" lang="en-US" '
        f'sz="{int(size*100)}" dirty="0"/>'))
    return p


def declare_a14_in_package(path):
    """保存後の pptx を開き直し、各スライドのルートに a14 名前空間を宣言する。

    PowerPoint が書き出すファイルと同じ形（xmlns:a14 と mc:Ignorable="a14" が
    p:sld にある）に揃えておく。数式を入れたスライドだけが対象。
    """
    import re
    import shutil
    import zipfile

    src = path + ".tmp"
    shutil.move(path, src)
    with zipfile.ZipFile(src) as zin, \
            zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if (item.filename.startswith("ppt/slides/slide")
                    and item.filename.endswith(".xml")
                    and b"a14:m" in data):
                text = data.decode("utf-8")
                text = re.sub(
                    r"<p:sld ([^>]*?)>",
                    lambda m: "<p:sld " + m.group(1)
                    + f' xmlns:a14="{A14_URI}" xmlns:mc="{MC_URI}"'
                      ' mc:Ignorable="a14">',
                    text, count=1)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    import os
    os.remove(src)
